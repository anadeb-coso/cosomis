# Dans MonApp/views.py

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from subprojects.models import Subproject, Project, SubprojectFile
from django.db.models import Q, Prefetch, Count
from django.utils.translation import gettext_lazy as _
from cosomis.constants import STRUCTURE_NOT_START_STATUS, STRUCTURE_IN_PROGRESS_ALL_STATUS, STRUCTURE_COMPLETED_STATUS, STRUCTURE_COMPLETED_ALL_STATUS, STRUCTURE_IN_PROGRESS_STATUS

def insufficient_subprojects_status_to_excel(request):
    
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    filename = "export_subprojects_images_status.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    # 2. Création du classeur et de la feuille de calcul
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Status des Images des ouvrages"

    # --- 3. Définition des En-têtes de Colonnes (Headers) ---
    
    columns_lazy = [
        _("Total"), _("Nombre d'ouvrages achevés"), _("Nombre d'ouvrages en cours"), _("Nombre d'ouvrages achevés sans 3 images"), _("Nombre d'ouvrages achevés avec 3 images"), _("Nombre d'ouvrages en cours sans images"), _("Nombre d'ouvrages en cours avec images"), 
    ]
    
    columns = [str(header) for header in columns_lazy]

    ws1.append(columns)

    project_name = request.GET.get("project_name", None)
    project  = Project.objects.filter(name=project_name).first()
    
    file_query = Q()
    for elt in STRUCTURE_COMPLETED_ALL_STATUS:
        file_query |= Q(subproject_step__wording__icontains=elt)
        file_query |= Q(name__icontains=elt)
        file_query |= Q(description__icontains=elt)
    file_query_for_subproject = Q()
    for elt in STRUCTURE_COMPLETED_ALL_STATUS:
        file_query_for_subproject |= Q(subprojectfile__subproject_step__wording__icontains=elt)
        file_query_for_subproject |= Q(subprojectfile__name__icontains=elt)
        file_query_for_subproject |= Q(subprojectfile__description__icontains=elt)
    
    subprojects = Subproject.objects.filter(
        current_status_of_the_site__in=STRUCTURE_COMPLETED_STATUS
    ).get_actifs().order_by("id").prefetch_related(
        Prefetch('subprojectfile_set', queryset=SubprojectFile.objects.filter(
            Q(
                Q(
                    subproject_step__wording__in=STRUCTURE_COMPLETED_STATUS,
                ) | 
                Q(
                    name__in=STRUCTURE_COMPLETED_STATUS,
                ) | 
                Q(
                    description__in=STRUCTURE_COMPLETED_STATUS,
                ) | 
                file_query
            )
        ).exclude(
            Q(url__icontains=".pdf") | Q(url__icontains=".doc")
        ), to_attr='photos')
    ).distinct()
    
    subprojects_completed_lt_3_images = Subproject.objects.filter(
        current_status_of_the_site__in=STRUCTURE_COMPLETED_STATUS
    ).get_actifs().annotate(
        photo_count=Count(
            'subprojectfile',
            filter=Q(
                Q(subprojectfile__subproject_step__wording__in=STRUCTURE_COMPLETED_STATUS) |
                Q(subprojectfile__name__in=STRUCTURE_COMPLETED_STATUS) |
                Q(subprojectfile__description__in=STRUCTURE_COMPLETED_STATUS) |
                file_query_for_subproject
            ) & ~Q(
                subprojectfile__url__icontains=".pdf"
            ) & ~Q(
                subprojectfile__url__icontains=".doc"
            ),
            distinct=True
        )
    ).filter(photo_count__lt=3)
    subprojects_completed_gte_3_images = Subproject.objects.filter(
        current_status_of_the_site__in=STRUCTURE_COMPLETED_STATUS
    ).get_actifs().annotate(
        photo_count=Count(
            'subprojectfile',
            filter=Q(
                Q(subprojectfile__subproject_step__wording__in=STRUCTURE_COMPLETED_STATUS) |
                Q(subprojectfile__name__in=STRUCTURE_COMPLETED_STATUS) |
                Q(subprojectfile__description__in=STRUCTURE_COMPLETED_STATUS) |
                file_query_for_subproject
            ) & ~Q(
                subprojectfile__url__icontains=".pdf"
            ) & ~Q(
                subprojectfile__url__icontains=".doc"
            ),
            distinct=True
        )
    ).filter(photo_count__gte=3)


    file_query_in_progress = Q()
    for elt in STRUCTURE_IN_PROGRESS_ALL_STATUS:
        file_query_in_progress |= Q(subprojectfile__subproject_step__wording__icontains=elt)
        file_query_in_progress |= Q(subprojectfile__name__icontains=elt)
        file_query_in_progress |= Q(subprojectfile__description__icontains=elt)

    subprojects_in_progress = (
        Subproject.objects
        .filter(current_status_of_the_site__in=STRUCTURE_IN_PROGRESS_STATUS)
        .get_actifs()
        .annotate(
            photo_count=Count(
                'subprojectfile',
                filter=Q(
                    Q(subprojectfile__subproject_step__wording__in=STRUCTURE_IN_PROGRESS_STATUS)
                    | file_query_in_progress
                ) & ~Q(subprojectfile__url__icontains=".pdf")
                & ~Q(subprojectfile__url__icontains=".doc"),
                distinct=True
            )
        ).distinct()
    )
    
    row_data = [
        # Total
        Subproject.objects.get_actifs().count(),

        # Nombre d'ouvrages achevés
        Subproject.objects.filter(
            current_status_of_the_site__in=STRUCTURE_COMPLETED_STATUS
        ).get_actifs().count(),

        # Nombre d'ouvrages en cours
        Subproject.objects.filter(
            current_status_of_the_site__in=STRUCTURE_IN_PROGRESS_STATUS
        ).get_actifs().count(),

        # Nombre d'ouvrages achevés sans 3 images
        subprojects_completed_lt_3_images.count(),

        # Nombre d'ouvrages achevés avec 3 images
        subprojects_completed_gte_3_images.count(),

        # Nombre d'ouvrages en cours sans images
        subprojects_in_progress.filter(photo_count=0).count(),

        # Nombre d'ouvrages en cours avec images
        subprojects_in_progress.filter(photo_count__gt=0).count(),
    ]
        
    ws1.append(row_data)

    # Ajustement de la largeur des colonnes (Optionnel) ---
    for col in range(1, len(columns) + 1):
        ws1.column_dimensions[get_column_letter(col)].width = 25



    # ================================== Details ===================================

    ws2 = wb.create_sheet(title="Villages_ouvrages sans 3 images")
    columns2 = ["Région", "Préfecture", "Commune", "Canton", "Village", "Type ouvrage", "Intitulé", "Composante"]
    ws2.append(columns2)
    _subprojects_completed_lt_3_images = subprojects_completed_lt_3_images.order_by(
        "location_subproject_realized__parent__parent__parent__parent__name",
        "location_subproject_realized__parent__parent__parent__name",
        "location_subproject_realized__parent__parent__name",
        "location_subproject_realized__parent__name",
        "location_subproject_realized__name"
    ).prefetch_related(
        'location_subproject_realized__parent__parent__parent__parent',
        'location_subproject_realized__parent__parent__parent',
        'location_subproject_realized__parent__parent',
        'location_subproject_realized__parent',
        'location_subproject_realized'
    )
    for subproject in _subprojects_completed_lt_3_images:
        ws2.append([
            subproject.location_subproject_realized.parent.parent.parent.parent.name,
            subproject.location_subproject_realized.parent.parent.parent.name,
            subproject.location_subproject_realized.parent.parent.name,
            subproject.location_subproject_realized.parent.name,
            subproject.location_subproject_realized.name,
            subproject.type_of_subproject,
            subproject.full_title_of_approved_subproject,
            subproject.component.name if subproject.component else "",
        ])

    for col in range(1, len(columns2) + 1):
        ws2.column_dimensions[get_column_letter(col)].width = 25

    # ================================== End - Details ===================================


    # ================================== No Coords ===================================

    ws3 = wb.create_sheet(title="Villages_ouvrages sans coords")
    columns3 = ["Région", "Préfecture", "Commune", "Canton", "Village", "Type ouvrage", "Intitulé", "Composante"]
    ws3.append(columns3)

    subprojects_completed_without_coords = Subproject.objects.filter(
        current_status_of_the_site__in=STRUCTURE_COMPLETED_STATUS
    ).filter(
        Q(latitude__isnull=True) | Q(latitude=0) | Q(latitude=0.0) | Q(longitude__isnull=True) | Q(longitude=0) | Q(longitude=0.0)
    ).get_actifs()

    _subprojects_completed_without_coords = subprojects_completed_without_coords.order_by(
        "location_subproject_realized__parent__parent__parent__parent__name",
        "location_subproject_realized__parent__parent__parent__name",
        "location_subproject_realized__parent__parent__name",
        "location_subproject_realized__parent__name",
        "location_subproject_realized__name"
    ).prefetch_related(
        'location_subproject_realized__parent__parent__parent__parent',
        'location_subproject_realized__parent__parent__parent',
        'location_subproject_realized__parent__parent',
        'location_subproject_realized__parent',
        'location_subproject_realized'
    )
    for subproject in _subprojects_completed_without_coords:
        ws3.append([
            subproject.location_subproject_realized.parent.parent.parent.parent.name,
            subproject.location_subproject_realized.parent.parent.parent.name,
            subproject.location_subproject_realized.parent.parent.name,
            subproject.location_subproject_realized.parent.name,
            subproject.location_subproject_realized.name,
            subproject.type_of_subproject,
            subproject.full_title_of_approved_subproject,
            subproject.component.name if subproject.component else "",
        ])

    for col in range(1, len(columns3) + 1):
        ws3.column_dimensions[get_column_letter(col)].width = 25

    # ================================== End - No Coords ===================================


    wb.save(response)
    
    return response


